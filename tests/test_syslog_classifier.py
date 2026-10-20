import os

import openpyxl
import pandas as pd
import pytest

import syslog_classifier as sc
from syslog_classifier import normalize_syslog as N


# ── normalize_syslog（README の置換表どおりか） ─────────────────
@pytest.mark.parametrize("src,expected", [
    ("2026-01-10 09:12:34 x", "<TS> x"),
    ("Jan 10 09:12:34 x", "<TS> x"),
    ("from 192.168.1.1 ok", "from <IP> ok"),
    ("addr fe80::1 dup", "addr <IPV> dup"),
    ("host aa:bb:cc:dd:ee:ff moved", "host <MAC> moved"),
    ("host aa-bb-cc-dd-ee-ff moved", "host <MAC> moved"),
    ("Link down on GigabitEthernet0/1", "Link down on <IF>"),
    ("value 0x1a2b", "value <HEX>"),
    ("kernel[12345]: oops", "<PROC> oops"),
    ("%SYS-5-CONFIG_I: Configured", "<SEV> Configured"),
    ("port 22 slot 3", "port <NUM> slot <NUM>"),
    ("router01 up", "router<NUM> up"),
])
def test_normalize_readme_table(src, expected):
    assert N(src) == expected


def test_normalize_iso_t_with_tz():
    assert N("2026-10-13T09:00:00+09:00 sw01 x").startswith("<TS>")


def test_normalize_rfc3164_single_digit_day():
    assert N("Oct  3 09:00:00 sw01 x").startswith("<TS>")


def test_normalize_idempotent_spaces():
    assert N("  a   b  ") == "a b"


# ── ローダー ─────────────────────────────────────────────
def test_load_text_filters(tmp_path):
    p = tmp_path / "a.log"
    p.write_text("# c\n\nline1\n   \nline2\r\n", encoding="utf-8")
    assert sc.load_new_logs_from_text(str(p)) == ["line1", "line2"]


def test_load_text_missing(tmp_path):
    with pytest.raises(FileNotFoundError):
        sc.load_new_logs_from_text(str(tmp_path / "nope.log"))


def test_load_known_missing_column(tmp_path):
    p = tmp_path / "l.xlsx"
    pd.DataFrame({"管理番号": ["A"]}).to_excel(p, sheet_name="台帳", index=False)
    with pytest.raises(ValueError):
        sc.load_known_logs(str(p))


def test_load_known_drops_nan(tmp_path):
    p = tmp_path / "l.xlsx"
    pd.DataFrame({"管理番号": ["A", "B", None], "ログ本文（raw）": ["x", None, "z"]}).to_excel(
        p, sheet_name="台帳", index=False)
    assert sc.load_known_logs(str(p)) == (["A"], ["x"])


def test_add_to_ledger(tmp_path):
    p = tmp_path / "l.xlsx"
    sc._create_sample_ledger(str(p))
    sc.add_to_ledger(str(p), "LOG-0101", "new log", "memo")
    ids, texts = sc.load_known_logs(str(p))
    assert ids[-1] == "LOG-0101" and texts[-1] == "new log" and len(ids) == 101
    with pytest.raises(ValueError):
        sc.add_to_ledger(str(p), "X", "y", sheet_name="nosheet")


# ── 検知器（モデルを使うのでモジュールで 1 回だけ作る） ─────────
@pytest.fixture(scope="module")
def det():
    return sc.UnknownDetector(
        resolved_logs=["2025-01-01 00:00:00 sw01 Link down on GigabitEthernet0/1",
                       "2025-01-01 00:00:00 rt01 BGP session with 10.0.0.1 reset by peer",
                       "2025-01-01 00:00:00 sw02 Fan failure detected in slot 2"],
        log_ids=["L1", "L2", "L3"])


def test_detect_empty(det):
    assert det.detect([]) == []


def test_detect_exact_template_is_known(det):
    r = det.detect(["2026-02-02 11:11:11 sw09 Link down on GigabitEthernet3/7"],
                   tau_high=0.85, tau_low=0.78, top_k=1)[0]
    assert r["decision"] == "known" and r["nearest_id"] == "L1" and r["score"] > 0.99


def test_detect_topk_larger_than_ledger(det):
    r = det.detect(["anything"], top_k=10)
    assert len(r) == 1


def test_detect_unrelated_is_strong(det):
    r = det.detect(["Kernel panic - not syncing: VFS unable to mount root fs"], top_k=1)[0]
    assert r["decision"] == "unknown_strong"


def test_check_threshold_returns_ordered(det):
    hi, lo = det.check_threshold()
    assert hi > lo


def test_cache_roundtrip_and_model_mismatch(det, tmp_path):
    p = str(tmp_path / "c.npz")
    det.save_cache(p)
    d2 = sc.UnknownDetector.load_cache(p)
    assert d2.log_ids == det.log_ids and d2.known_normed == det.known_normed
    a = det.detect(["sw01 Fan failure detected in slot 1"])[0]["score"]
    b = d2.detect(["sw01 Fan failure detected in slot 1"])[0]["score"]
    assert a == pytest.approx(b, abs=1e-4)
    with pytest.raises(ValueError):
        sc.UnknownDetector.load_cache(p, model_name="other-model")


def test_update_known(det, tmp_path):
    n = len(det.log_ids)
    det.update_known("sw01 Kernel panic - not syncing", "L9")
    assert len(det.log_ids) == n + 1 and det.known_vecs.shape[0] == n + 1
    r = det.detect(["sw05 Kernel panic - not syncing"], top_k=1)[0]
    assert r["nearest_id"] == "L9" and r["decision"] == "known"


def test_cluster_and_export(det, tmp_path):
    logs = ["sw01 Kernel oops in module foo"] * 3 + ["sw01 Link down on GigabitEthernet0/1"] + \
           ["UPS switched to battery power"] * 2
    res = sc.cluster_unknowns(det.detect(logs, tau_high=0.85, tau_low=0.78))
    unk = [r for r in res if r["decision"] != "known"]
    assert all("_vec" not in r for r in res)
    assert sum(r["is_representative"] for r in unk) == len({r["cluster_id"] for r in unk})
    out = sc.export_unknowns_to_excel(res, str(tmp_path / "o.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["クラスタサマリー", "要確認ログ（詳細）"]


def test_cluster_all_known(det):
    res = sc.cluster_unknowns(det.detect(["sw03 Link down on GigabitEthernet1/1"], tau_high=0.5, tau_low=0.4))
    assert res[0]["cluster_id"] == -1


def test_export_without_cluster(det, tmp_path):
    res = det.detect(["totally new event xyz"], tau_high=0.99, tau_low=0.98)
    for r in res:
        r.pop("_vec")
    wb = openpyxl.load_workbook(sc.export_unknowns_to_excel(res, str(tmp_path / "o.xlsx")))
    assert wb.sheetnames == ["要確認ログ"]


def test_export_control_chars(det, tmp_path):
    """実際の syslog にはエスケープシーケンス等の制御文字が混入することがある"""
    res = sc.cluster_unknowns(det.detect(["sw01 weird \x1b[31mcolor\x1b[0m output \x00"],
                                         tau_high=0.99, tau_low=0.98))
    wb = openpyxl.load_workbook(sc.export_unknowns_to_excel(res, str(tmp_path / "o.xlsx")))
    assert wb["要確認ログ（詳細）"].cell(2, 8).value == "sw01 weird [31mcolor[0m output "


def test_normalize_mac_in_message():
    assert N("Host 00:1a:2b:3c:4d:5e in vlan 10 is flapping") == "Host <MAC> in vlan <NUM> is flapping"


def test_print_summary_empty(capsys):
    sc.print_summary([])
    assert "空" in capsys.readouterr().out
